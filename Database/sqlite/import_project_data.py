#!/usr/bin/env python3
"""Import Access-export CSV/XLSX files into the project SQLite database."""

from __future__ import annotations

import argparse
import csv
import json
import sqlite3
import sys
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


DEF_TABLE_MAP = {
    "Ballinacurra": "ballinacurra",
    "BurialCertificates": "burial_certificate",
    "BurialPlots": "burial_plot",
    "Cemetery": "cemetery",
    "Data": "data",
    "Headstones": "headstone",
    "Interred": "interred",
    "IRAMedalsApplicants": "iramedals_applicants",
    "IRAMedalsDetails": "iramedals_details",
    "InterredStDavids": "interred_st_davids",
    "MaritalStatus": "marital_status",
    "Months": "month_lookup",
    "Relationships": "relationship_lookup",
    "ReligiousDenominations": "religious_denomination",
    "StFinbarrs": "st_finbarrs",
    "WebsiteExportCivil": "website_export_civil",
    "WebsiteExportInterred": "website_export_interred",
    "StandardChristianNames": "standard_christian_names",
    "SurnamesStandardList": "surnames_standard_list",
    "Townlands": "townland",
}

SOURCE_COLUMN_ALIASES = {
    "iramedals_applicants": {
        "medalid": "medal_id",
        "townlandid": "townland_id",
        "nameid": "name_id",
        "surnameid": "surname_id",
        "fileref": "file_ref",
        "mob": "mob",
        "associatedpensionfile": "associated_pension_file",
    },
    "iramedals_details": {
        "medalid": "medal_id",
        "maidenid": "maiden_id",
        "statusid": "status_id",
        "successfulmedal": "successful_medal",
        "medalawarded": "medal_awarded",
        "pensioned": "pensioned",
        "pensionrejectedother": "pension_rejected_other",
        "organisationid": "organisation_id",
    },
    "cemetery": {"cemeteryid": "cemetery_id", "cemetery": "name"},
    "townland": {"townlandid": "townland_id", "townland": "townland_name"},
    "standard_christian_names": {"nameid": "christian_name_id", "name": "christian_name"},
    "surnames_standard_list": {"surnameid": "surname_id", "surname": "surname"},
    "burial_certificate": {
        "certid": "burial_certificate_id",
        "burialid": "burial_id",
        "day": "certificate_day",
        "certmonthid": "certificate_month_id",
        "year": "certificate_year",
        "age": "age_at_certificate",
        "statusid": "status_id",
        "occupation": "occupation",
        "address": "address",
        "townlandid": "townland_id",
        "townland": "townland_name",
        "witness": "witness",
        "witrelationship": "witness_relationship",
        "conmments": "comments",
        "comments": "comments",
    },
    "burial_plot": {
        "burialid": "burial_id",
        "cemeteryid": "cemetery_id",
        "plot": "plot_number",
        "row": "row_number",
        "inscription": "inscription",
        "headstone": "headstone",
        "comments": "comments",
        "ownersurnameid": "owner_surname_id",
        "ownersurname": "owner_surname",
        "ownerchristianid": "owner_christian_id",
        "ownerchristian": "owner_christian",
        "owneralias": "owner_alias",
        "townland": "townland_name",
        "relationshipid": "relationship_id",
    },
    "interred": {
        "interredid": "access_interred_id",
        "interedid": "access_interred_id",
        "no": "record_number",
        "burialid": "burial_id",
        "interredrelationshipid": "relationship_id",
        "interedrelationshipid": "relationship_id",
        "interredtownlandid": "townland_id",
        "interedtownlandid": "townland_id",
        "christianid": "christian_name_id",
        "christianname": "first_name",
        "age": "age_at_death",
        "born": "date_of_birth",
        "dayintered": "burial_day",
        "monthintered": "burial_month",
        "yearintered": "burial_year",
        "intersurnameid": "surname_id",
        "interredsurnameid": "surname_id",
        "interedsurnameid": "surname_id",
        "plot": "plot_number",
        "row": "row_number",
    },
    "relationship_lookup": {"relationshipid": "relationship_id", "relationship": "relationship_name"},
    "marital_status": {"maritalstatusid": "marital_status_id", "status": "status_name"},
    "religious_denomination": {"religiousdenominationid": "religion_id", "religiousdenomination": "denomination_name"},
}


def normalize_identifier(value: str) -> str:
    value = value.strip()
    value = value.replace("-", "_")
    value = value.replace(" ", "_")
    value = value.replace("/", "_")
    cleaned = []
    for ch in value:
        if ch.isalnum() or ch == "_":
            cleaned.append(ch)
        else:
            cleaned.append("_")
    text = "".join(cleaned)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_").lower()


def read_json_map(path: Path):
    if not path.exists():
        return DEF_TABLE_MAP.copy()
    with path.open("r", encoding="utf-8") as fh:
        data = json.load(fh)
    return {str(k): str(v) for k, v in data.items()}


def table_name_from_file(file_path: Path, mapping: dict[str, str]) -> str | None:
    stem = file_path.stem
    if stem in mapping:
        return mapping[stem]
    normalized = normalize_identifier(stem)
    for source_name, target_name in mapping.items():
        if normalize_identifier(source_name) == normalized:
            return target_name
    for source_name, target_name in mapping.items():
        if normalize_identifier(source_name).replace("_", "") == normalized.replace("_", ""):
            return target_name
    return None


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        rows = []
        for row in reader:
            cleaned = {}
            for key, value in row.items():
                if key is None:
                    continue
                cleaned[key] = "" if value is None else value.strip()
            rows.append(cleaned)
    return rows


def read_xlsx_rows(path: Path, sheet_name: str | None = None) -> list[dict[str, str]]:
    if load_workbook is None:
        raise RuntimeError(
            f"openpyxl is required to read Excel files. Install it with: pip install openpyxl"
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    rows = []
    first = True
    header = []
    for row in sheet.iter_rows(values_only=True):
        if first:
            header = [str(cell).strip() if cell is not None else "" for cell in row]
            first = False
            continue
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue
        item = {}
        for idx, cell in enumerate(row):
            if idx >= len(header):
                break
            value = "" if cell is None else str(cell).strip()
            if str(header[idx]).casefold() == "address":
                value = value.replace("_x000d_", "\n").replace("_x000a_", "\n")
                item[header[idx]] = "\n".join(line.strip() for line in value.splitlines() if line.strip())
            else:
                value = value.replace("_x000d_", " ").replace("_x000a_", " ")
                item[header[idx]] = " ".join(value.split())
        rows.append(item)
    workbook.close()
    return rows


def xlsx_sheet_names(path: Path) -> list[str]:
    if load_workbook is None:
        raise RuntimeError(
            "openpyxl is required to read Excel files. Install it with: pip install openpyxl"
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    names = workbook.sheetnames
    workbook.close()
    return names


def pragma_columns(conn: sqlite3.Connection, table_name: str) -> list[str]:
    info = conn.execute(f'PRAGMA table_info("{table_name}")').fetchall()
    return [row[1] for row in info]


def primary_key_columns(conn: sqlite3.Connection, table_name: str) -> list[str]:
    info = conn.execute(f'PRAGMA table_info("{table_name}")').fetchall()
    return [row[1] for row in info if row[5]]


def import_month_lookup(conn: sqlite3.Connection, source_db: Path) -> int:
    if not source_db.exists():
        return 0
    source = sqlite3.connect(f"file:{source_db}?mode=ro", uri=True)
    try:
        rows = source.execute("SELECT MonthsID, Months FROM Months ORDER BY MonthsID").fetchall()
    except sqlite3.Error:
        source.close()
        return 0
    source.close()
    for month_id, month_name in rows:
        conn.execute(
            "INSERT OR REPLACE INTO month_lookup (month_id, month_number, month_name) VALUES (?, ?, ?)",
            (month_id, month_id, month_name),
        )
    conn.commit()
    return len(rows)


def apply_address_amendments(conn: sqlite3.Connection, workbook_path: Path) -> int:
    if not workbook_path.exists() or load_workbook is None:
        return 0
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"]
        rows = read_xlsx_rows(workbook_path, "Sheet1")
    except (KeyError, ValueError):
        workbook.close()
        return 0
    workbook.close()
    updated = 0
    for row in rows:
        medal_id = row.get("MedalID")
        address = row.get("Address")
        if medal_id in (None, ""):
            continue
        conn.execute(
            "UPDATE iramedals_applicants SET townland_id = COALESCE(?, townland_id), address = ? WHERE medal_id = ?",
            (row.get("TownlandId") or None, address or None, int(float(medal_id))),
        )
        updated += conn.execute("SELECT changes()").fetchone()[0]
    conn.commit()
    return updated


def record_source_file(conn: sqlite3.Connection, file_path: Path) -> int:
    original_name = file_path.name
    conn.execute(
        "INSERT INTO source_file (original_name, source_type) VALUES (?, ?)",
        (original_name, "csv" if file_path.suffix.lower() == ".csv" else "excel"),
    )
    return int(conn.execute("SELECT last_insert_rowid()").fetchone()[0])


def import_file(conn: sqlite3.Connection, file_path: Path, mapping: dict[str, str]) -> tuple[str, int]:
    if file_path.suffix.lower() == ".csv":
        table_name = table_name_from_file(file_path, mapping)
        if table_name is None:
            return ("SKIPPED", 0)
        rows = read_csv_rows(file_path)
    elif file_path.suffix.lower() in {".xlsx", ".xls"}:
        sheet_names = xlsx_sheet_names(file_path)
        imported = []
        for sheet_name in sheet_names:
            table_name = table_name_from_source_name(sheet_name, mapping)
            if table_name is None:
                continue
            count = import_rows(conn, table_name, read_xlsx_rows(file_path, sheet_name), file_path)
            imported.append((table_name, count))
        if not imported:
            return ("SKIPPED", 0)
        return (", ".join(table_name for table_name, _ in imported), sum(count for _, count in imported))
    else:
        return ("SKIPPED", 0)

    return table_name, import_rows(conn, table_name, rows, file_path)


def table_name_from_source_name(source_name: str, mapping: dict[str, str]) -> str | None:
    normalized = normalize_identifier(source_name)
    for name, target_name in mapping.items():
        if normalize_identifier(name) == normalized:
            return target_name
    return None


def import_rows(conn: sqlite3.Connection, table_name: str, rows: list[dict[str, str]], file_path: Path) -> int:
    source_file_id = record_source_file(conn, file_path)
    columns = pragma_columns(conn, table_name)
    primary_keys = primary_key_columns(conn, table_name)
    normalized_columns = {normalize_identifier(col): col for col in columns}
    for source_name, target_name in SOURCE_COLUMN_ALIASES.get(table_name, {}).items():
        if target_name in columns:
            normalized_columns[source_name] = target_name

    inserted = 0
    for row in rows:
        if not row:
            continue
        normalized_row = {}
        for raw_key, value in row.items():
            key = normalize_identifier(str(raw_key))
            if key in normalized_columns:
                normalized_row[normalized_columns[key]] = value if value != "" else None
        if table_name == "interred":
            source_values = {
                normalize_identifier(str(raw_key)): value
                for raw_key, value in row.items()
            }
            day = source_values.get("dayintered")
            month = source_values.get("monthintered")
            year = source_values.get("yearintered")
            for source_key, target_key in (
                ("dayintered", "burial_day"),
                ("monthintered", "burial_month"),
                ("yearintered", "burial_year"),
            ):
                value = source_values.get(source_key)
                if value not in (None, "") and target_key in columns:
                    try:
                        normalized_row[target_key] = int(float(value))
                    except (TypeError, ValueError):
                        pass
            if year not in (None, "") and "burial_date" in columns:
                try:
                    year_value = int(float(year))
                    month_value = int(float(month)) if month not in (None, "") else None
                    day_value = int(float(day)) if day not in (None, "") else None
                    if month_value and day_value:
                        normalized_row["burial_date"] = f"{year_value:04d}-{month_value:02d}-{day_value:02d}"
                    elif month_value:
                        normalized_row["burial_date"] = f"{year_value:04d}-{month_value:02d}"
                    else:
                        normalized_row["burial_date"] = str(year_value)
                except (TypeError, ValueError):
                    pass
        if table_name == "iramedals_applicants":
            source_values = {
                normalize_identifier(str(raw_key)): value
                for raw_key, value in row.items()
            }
            death = source_values.get("death")
            if death not in (None, ""):
                text = str(death).strip()
                parts = text[:10].split("-")
                if len(parts) == 3 and all(part.strip().isdigit() for part in parts):
                    try:
                        year, month, day = (int(part.strip()) for part in parts)
                        normalized_row["death_day"] = day
                        normalized_row["death_month"] = month
                        normalized_row["death_year"] = year
                    except ValueError:
                        pass
        if "source_file_id" in normalized_columns and "source_file_id" not in normalized_row:
            normalized_row["source_file_id"] = source_file_id
        if not normalized_row:
            continue

        insert_cols = []
        insert_values = []
        for col_name in columns:
            if col_name.endswith("_id") and col_name not in normalized_row:
                continue
            if col_name in normalized_row:
                insert_cols.append(col_name)
                insert_values.append(normalized_row[col_name])

        if not insert_cols:
            continue

        placeholders = ", ".join(["?"] * len(insert_cols))
        sql = f'INSERT OR IGNORE INTO "{table_name}" ({", ".join(f'"{c}"' for c in insert_cols)}) VALUES ({placeholders})'
        conn.execute(sql, insert_values)
        if primary_keys and all(key in normalized_row for key in primary_keys):
            update_cols = [col for col in insert_cols if col not in primary_keys]
            if update_cols:
                assignments = ", ".join(f'"{col}" = ?' for col in update_cols)
                where = " AND ".join(f'"{key}" = ?' for key in primary_keys)
                values = [normalized_row[col] for col in update_cols]
                values.extend(normalized_row[key] for key in primary_keys)
                conn.execute(
                    f'UPDATE "{table_name}" SET {assignments} WHERE {where}',
                    values,
                )
        inserted += 1

    conn.commit()
    return inserted


def iter_data_files(raw_dir: Path) -> Iterable[Path]:
    if not raw_dir.exists():
        return []
    files = []
    for pattern in ("*.csv", "*.xlsx", "*.xls"):
        files.extend(raw_dir.rglob(pattern))
    return sorted(set(files))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Import Access export files into the SQLite project database.")
    parser.add_argument("--raw-dir", type=Path, default=Path(__file__).resolve().parent.parent / "raw")
    parser.add_argument("--db", type=Path, default=Path(__file__).resolve().parent / "project.db")
    parser.add_argument("--schema", type=Path, default=Path(__file__).resolve().parent / "schema.sql")
    parser.add_argument("--map", type=Path, default=Path(__file__).resolve().parent / "table_map.json")
    parser.add_argument("--reset", action="store_true", help="Delete the existing database before import.")
    parser.add_argument("--extra-file", type=Path, action="append", default=[], help="Additional CSV/XLSX/XLS source file to import.")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    db_path = args.db
    if args.reset and db_path.exists():
        db_path.unlink()

    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(db_path))
    conn.execute("PRAGMA foreign_keys = ON")

    if args.schema.exists():
        schema_sql = args.schema.read_text(encoding="utf-8")
        conn.executescript(schema_sql)
    else:
        print(f"Schema not found: {args.schema}", file=sys.stderr)
        return 1

    mapping = read_json_map(args.map)
    files = list(iter_data_files(args.raw_dir))
    files.extend(path for path in args.extra_file if path.exists())
    if not files:
        print(f"No CSV/XLSX/XLS files found in {args.raw_dir}")
        return 0

    imported = []
    for file_path in files:
        table_name, count = import_file(conn, file_path, mapping)
        if table_name != "SKIPPED":
            for target_name in table_name.split(", "):
                imported.append((target_name, count, file_path.name))
            print(f"{file_path.name} -> {table_name}: {count} rows imported")

    months_source = args.db.parent.parent / "Carrigtwohill.db"
    months_imported = import_month_lookup(conn, months_source)
    if months_imported:
        print(f"{months_source.name} -> month_lookup: {months_imported} rows imported")

    amendments = args.db.parent.parent / "queries" / "Address Amendments.xlsx"
    amendments_applied = apply_address_amendments(conn, amendments)
    if amendments_applied:
        print(f"{amendments.name} -> iramedals_applicants: {amendments_applied} addresses amended")

    print("\nValidation summary:")
    for table_name, count, file_name in imported:
        actual = conn.execute(f'SELECT COUNT(*) FROM "{table_name}"').fetchone()[0]
        print(f"- {table_name}: {actual} rows in SQLite ({file_name})")

    conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
